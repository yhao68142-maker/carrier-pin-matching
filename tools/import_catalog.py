"""Read-only extraction; never alters the source workbook."""
import json, re, sys
from pathlib import Path
import openpyxl

def extract(source):
    workbook=openpyxl.load_workbook(source,data_only=True)
    records=[]
    def merged_value(sheet,row,col):
        cell=sheet.cell(row,col)
        for merged in sheet.merged_cells.ranges:
            if cell.coordinate in merged:
                cell=sheet.cell(merged.min_row,merged.min_col)
                break
        return cell.value,cell.coordinate
    outlines=[]
    def pin_pair(raw):
        found=re.match(r'\s*(\d+(?:\.\d+)?)\s*[*×xX]\s*(\d+(?:\.\d+)?)',str(raw or ''))
        return tuple(found.groups()) if found else None
    cb=workbook['治具料盒信息统计（CB端）']
    for row in range(1,cb.max_row+1):
        if cb.cell(row,4).value!='载条定位销尺寸': continue
        note,cell=merged_value(cb,row,6)
        found=re.search(r'载条外形尺寸[：:]\s*(\d+(?:\.\d+)?)\s*[*×xX]\s*(\d+(?:\.\d+)?)',str(note or ''))
        if found:
            outlines.append({'code':str(merged_value(cb,row,2)[0] or '').strip(),
                             'pins':pin_pair(cb.cell(row,5).value),
                             'length':found[1],'width':found[2],
                             'outerSheet':cb.title,'outerCell':cell,'outerRaw':found[0]})
    for sheet in workbook:
        if sheet.title not in ('治具料盒信息统计','治具料盒信息统计（CB端）'): continue
        def value(row,col):
            cell=sheet.cell(row,col)
            for merged in sheet.merged_cells.ranges:
                if cell.coordinate in merged:
                    return sheet.cell(merged.min_row,merged.min_col).value
            return cell.value
        for row in range(1,sheet.max_row+1):
            if sheet.cell(row,4).value != '载条定位销尺寸': continue
            raw=str(sheet.cell(row,5).value or '')
            found=re.match(r'\s*(\d+(?:\.\d+)?)\s*[*×xX]\s*(\d+(?:\.\d+)?)',raw)
            if not found: continue
            record={'code':str(value(row,2) or '').strip(),'model':str(value(row,3) or ''),
                'x':found[1],'y':found[2],'raw':raw,'sheet':sheet.title,'cell':f'E{row}'}
            candidates=[o for o in outlines if record['code'] and o['code']==record['code'] and o['pins']==pin_pair(raw)]
            if sheet.title==cb.title:
                note,note_cell=merged_value(sheet,row,6)
                outer=re.search(r'载条外形尺寸[：:]\s*(\d+(?:\.\d+)?)\s*[*×xX]\s*(\d+(?:\.\d+)?)',str(note or ''))
                candidates=[{'length':outer[1],'width':outer[2],'outerSheet':sheet.title,
                             'outerCell':note_cell,'outerRaw':outer[0]}] if outer else []
            if len({(o['length'],o['width']) for o in candidates})==1:
                record.update({k:v for k,v in candidates[0].items() if k not in ('code','pins')})
            else:
                record.update(length=None,width=None,outerSheet=None,outerCell=None,outerRaw=None)
            records.append(record)
    return records

if __name__=='__main__':
    dest=Path(__file__).resolve().parents[1]/'data'/'catalog.json'
    dest.parent.mkdir(exist_ok=True)
    rows=extract(sys.argv[1])
    dest.write_text(json.dumps(rows,ensure_ascii=False,indent=2),encoding='utf-8')
    print(f'Extracted {len(rows)} records; original workbook unchanged.')
